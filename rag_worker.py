"""CPU OCR, embeddings and sharded FAISS indexing for the local vault."""
import json
import os
from pathlib import Path
import re
import secrets
import sqlite3
import sys
import time

import faiss
import numpy as np
from fastembed import TextEmbedding
from fastembed.rerank.cross_encoder import TextCrossEncoder

PROJECT = Path(__file__).resolve().parent
ROOT = Path(os.environ.get('AIWS_KNOWLEDGE_ROOT', PROJECT / 'data'))
DB = ROOT / 'knowledge.sqlite3'
INDEX_DIR = ROOT / 'faiss'
FASTEMBED_CACHE = os.environ.get('AIWS_FASTEMBED_CACHE', str(PROJECT / '.cache/fastembed'))
DIM = 384
SHARD_LIMIT = 50000
MODEL = None
RERANKER = None
OCR = None


def connect():
    con = sqlite3.connect(DB, timeout=60)
    con.row_factory = sqlite3.Row
    con.execute('PRAGMA busy_timeout=60000')
    return con


def model():
    global MODEL
    if MODEL is None:
        MODEL = TextEmbedding(model_name='BAAI/bge-small-en-v1.5', cache_dir=FASTEMBED_CACHE)
    return MODEL


def reranker():
    global RERANKER
    if RERANKER is None:
        RERANKER = TextCrossEncoder(model_name='jinaai/jina-reranker-v1-turbo-en',
                                   cache_dir=FASTEMBED_CACHE,threads=4)
    return RERANKER


def ocr(image):
    global OCR
    if OCR is None:
        from rapidocr_onnxruntime import RapidOCR
        OCR = RapidOCR()
    found, _ = OCR(image)
    return '\n'.join(item[1] for item in (found or []) if len(item) > 1)


def pages(path):
    suffix = path.suffix.lower()
    if suffix == '.pdf':
        import fitz
        document = fitz.open(path)
        for number, page in enumerate(document, 1):
            text = page.get_text(sort=True).strip()
            used_ocr = False
            if len(text) < 60:
                pix = page.get_pixmap(matrix=fitz.Matrix(1.6,1.6), alpha=False)
                image = np.frombuffer(pix.samples,dtype=np.uint8).reshape(pix.height,pix.width,pix.n)
                text = ocr(image)
                used_ocr = True
            yield number,text,used_ocr
    elif suffix in {'.png','.jpg','.jpeg','.webp','.tif','.tiff'}:
        from PIL import Image
        Image.MAX_IMAGE_PIXELS = 100_000_000
        with Image.open(path) as image:
            frames = getattr(image,'n_frames',1)
            for number in range(frames):
                image.seek(number)
                yield number+1,ocr(np.asarray(image.convert('RGB'))),True
    elif suffix == '.docx':
        from docx import Document
        document = Document(path)
        text = '\n'.join(p.text for p in document.paragraphs)
        for table in document.tables:
            text += '\n' + '\n'.join(' | '.join(c.text for c in row.cells) for row in table.rows)
        yield 1,text,False
    else:
        with path.open('r',encoding='utf-8-sig',errors='replace') as file:
            page = 1
            while block := file.read(50000):
                yield page,block,False
                page += 1


def chunks(text):
    text = re.sub(r'\x00','',text).strip()
    start = 0
    while start < len(text):
        end = min(len(text),start+1500)
        if end < len(text):
            cut = max(text.rfind('\n',start+800,end),text.rfind(' ',start+800,end))
            if cut > start+800:
                end = cut
        part = text[start:end].strip()
        if part:
            yield part
        if end == len(text):
            break
        start = max(start+1,end-180)


def index_file(shard):
    return INDEX_DIR / f'index-{shard:05d}.faiss'


def new_index():
    base = faiss.IndexHNSWFlat(DIM,32,faiss.METRIC_INNER_PRODUCT)
    base.hnsw.efConstruction = 80
    base.hnsw.efSearch = 64
    return faiss.IndexIDMap2(base)


def load_index(con, shard):
    INDEX_DIR.mkdir(parents=True,exist_ok=True)
    path = index_file(shard)
    index = faiss.read_index(str(path)) if path.exists() else new_index()
    count = con.execute('SELECT COUNT(*) FROM chunks WHERE shard=?',(shard,)).fetchone()[0]
    if index.ntotal != count:
        index = new_index()
        cursor = con.execute('SELECT id,embedding FROM chunks WHERE shard=? ORDER BY id',(shard,))
        while rows := cursor.fetchmany(512):
            vectors = np.vstack([np.frombuffer(row['embedding'],dtype=np.float32) for row in rows])
            ids = np.array([row['id'] for row in rows],dtype=np.int64)
            index.add_with_ids(vectors,ids)
        save_index(index,shard)
    return index


def save_index(index,shard):
    path = index_file(shard)
    temp = path.with_name(f'{path.name}.{os.getpid()}-{secrets.token_hex(4)}.tmp')
    try:
        faiss.write_index(index,str(temp))
        temp.replace(path)
    finally:
        temp.unlink(missing_ok=True)


def ingest(document_id):
    con = connect()
    row = con.execute('SELECT * FROM documents WHERE id=?',(document_id,)).fetchone()
    if row is None:
        raise ValueError('Document not found')
    if row['status'] == 'ready':
        return {'id':document_id,'status':'ready'}
    old_shards = [r[0] for r in con.execute('SELECT DISTINCT shard FROM chunks WHERE document_id=?',(document_id,))]
    if old_shards:
        con.execute('DELETE FROM chunks_fts WHERE rowid IN (SELECT id FROM chunks WHERE document_id=?)',(document_id,))
        con.execute('DELETE FROM chunks WHERE document_id=?',(document_id,))
        con.execute('UPDATE documents SET chunk_count=0,ocr_pages=0 WHERE id=?',(document_id,))
        con.commit()
        for shard in old_shards:
            load_index(con,shard)
    con.execute("UPDATE documents SET status='indexing',error=NULL,updated=? WHERE id=?",(time.time(),document_id))
    con.commit()
    current_shard = con.execute('SELECT COALESCE(MAX(shard),0) FROM chunks').fetchone()[0]
    index = load_index(con,current_shard)
    pending = []
    count = con.execute('SELECT COUNT(*) FROM chunks WHERE document_id=?',(document_id,)).fetchone()[0]
    ocr_pages = 0

    def flush():
        nonlocal current_shard,index,count,pending
        if not pending:
            return
        vectors = np.array(list(model().embed([item[1] for item in pending],batch_size=32)),dtype=np.float32)
        faiss.normalize_L2(vectors)
        cursor = con.cursor()
        for (page,text), vector in zip(pending,vectors):
            if index.ntotal >= SHARD_LIMIT:
                save_index(index,current_shard)
                current_shard += 1
                index = load_index(con,current_shard)
            cursor.execute('INSERT INTO chunks(document_id,page,text,embedding,shard) VALUES(?,?,?,?,?)',
                           (document_id,page,text,vector.tobytes(),current_shard))
            chunk_id = cursor.lastrowid
            cursor.execute('INSERT INTO chunks_fts(rowid,text) VALUES(?,?)',(chunk_id,text))
            index.add_with_ids(vector.reshape(1,-1),np.array([chunk_id],dtype=np.int64))
            count += 1
        con.execute('UPDATE documents SET chunk_count=?,ocr_pages=?,updated=? WHERE id=?',
                    (count,ocr_pages,time.time(),document_id))
        con.commit()
        save_index(index,current_shard)
        pending = []

    try:
        for page,text,used_ocr in pages(Path(row['path'])):
            ocr_pages += int(used_ocr)
            for part in chunks(text):
                pending.append((page,part))
                if len(pending) >= 64:
                    flush()
        flush()
        if not count:
            raise ValueError('No readable text or OCR content found in this file.')
        con.execute("UPDATE documents SET status='ready',chunk_count=?,ocr_pages=?,updated=? WHERE id=?",
                    (count,ocr_pages,time.time(),document_id))
        con.commit()
        return {'id':document_id,'status':'ready','chunks':count,'ocr_pages':ocr_pages}
    finally:
        con.close()


def search(question,limit):
    con = connect()
    try:
        vector = np.asarray(next(model().query_embed(question)),dtype=np.float32).reshape(1,-1)
        faiss.normalize_L2(vector)
        scores = {}
        for path in sorted(INDEX_DIR.glob('index-*.faiss')) if INDEX_DIR.exists() else []:
            shard = int(path.stem.split('-')[1])
            index = load_index(con,shard)
            similarities, ids = index.search(vector,max(30,limit*5))
            for similarity, chunk_id in zip(similarities[0],ids[0]):
                if chunk_id >= 0:
                    scores[int(chunk_id)] = max(scores.get(int(chunk_id),0),float(similarity))
        terms = re.findall(r'\w+',question,flags=re.UNICODE)[:12]
        if terms:
            query = ' OR '.join('"'+term.replace('"','')+'"' for term in terms)
            try:
                for row in con.execute('SELECT rowid FROM chunks_fts WHERE chunks_fts MATCH ? LIMIT 100',(query,)):
                    scores[row['rowid']] = scores.get(row['rowid'],0) + .12
            except sqlite3.OperationalError:
                pass
        ranked = sorted(scores,key=lambda key:scores[key],reverse=True)
        candidates = []
        for chunk_id in ranked[:max(24,limit*5)]:
            row = con.execute("SELECT c.id,c.page,c.text,d.id document_id,d.name FROM chunks c JOIN documents d ON d.id=c.document_id WHERE c.id=? AND d.status='ready'",(chunk_id,)).fetchone()
            if row:
                candidates.append({'chunk_id':row['id'],'document_id':row['document_id'],
                                   'name':row['name'],'page':row['page'],'text':row['text'],
                                   'score':round(scores[chunk_id],3)})
        ranking = 'hybrid'
        if candidates:
            try:
                rerank_scores = list(reranker().rerank(question,[item['text'] for item in candidates],batch_size=16))
                if len(rerank_scores) != len(candidates):
                    raise RuntimeError('Reranker returned an incomplete score list.')
                for item, score in zip(candidates,rerank_scores):
                    item['rerank_score'] = round(float(score),3)
                candidates.sort(key=lambda item:item['rerank_score'],reverse=True)
                ranking = 'cross-encoder'
            except Exception:
                import traceback
                traceback.print_exc(file=sys.stderr)
                ranking = 'hybrid-fallback'
        diverse = []
        per_document = {}
        for item in candidates:
            if per_document.get(item['document_id'],0) < 3:
                diverse.append(item)
                per_document[item['document_id']] = per_document.get(item['document_id'],0)+1
            if len(diverse) >= limit:
                break
        if len(diverse) < limit:
            seen = {item['chunk_id'] for item in diverse}
            diverse.extend(item for item in candidates if item['chunk_id'] not in seen)
        return {'sources':diverse[:limit],'ranking':ranking}
    finally:
        con.close()


if __name__ == '__main__':
    try:
        if hasattr(sys.stdout,'reconfigure'):
            sys.stdout.reconfigure(encoding='utf-8')
        if sys.argv[1] == 'ingest':
            result = ingest(sys.argv[2])
        elif sys.argv[1] == 'search':
            result = search(sys.argv[2],min(20,max(1,int(sys.argv[3]))))
        else:
            raise ValueError('Unknown worker command')
        print(json.dumps(result,ensure_ascii=False))
    except Exception as exc:
        import traceback
        traceback.print_exc(file=sys.stderr)
        sys.exit(1)
